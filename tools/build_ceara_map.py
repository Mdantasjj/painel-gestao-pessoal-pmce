from __future__ import annotations

import csv
import json
import re
import unicodedata
import urllib.request
import zipfile
from pathlib import Path

import shapefile
from openpyxl import load_workbook
from pyproj import Transformer
from shapely.geometry import MultiPolygon, Polygon, mapping, shape
from shapely.ops import transform, unary_union


ROOT = Path(__file__).resolve().parents[1]
MAP_DATA = ROOT / "mapa" / "data"
IPECE_DIR = ROOT / "data" / "ipece_limites_2026"
IPECE_ARCHIVE = ROOT / "data" / "Limites_municipais_Ceara_2026.zip"
IPECE_URL = "https://www.ipece.ce.gov.br/wp-content/uploads/sites/45/2026/05/Limites_municipais_Ceara_2026.zip"
DISTRI_VTR_CANDIDATES = (
    ROOT / "PACOTE_FONTES_DADOS_PMCE_2026" / "01_FONTES_ORIGINAIS_USADAS" / "DISTRI VTR (1).xlsx",
    Path.home() / "Downloads" / "DISTRI VTR (1).xlsx",
)
FORTALEZA_SOURCE = MAP_DATA / "bairros-fortaleza-2023.geojson"
FORTALEZA_URL = "https://mapas.fortaleza.ce.gov.br/api/download/geojson/21"
CSV_SOURCE = ROOT / "data" / "saidas_batalhoes_2026.csv"


COMPANY_STRENGTHS = {
    "1º BPM": [178, 116], "2º BPM": [305, 121, 160, 1], "3º BPM": [112, 96, 81, 117],
    "4º BPM": [159, 81], "5º BPM": [245, 63, 77], "6º BPM": [126, 105, 109],
    "7º BPM": [150, 73, 89], "8º BPM": [229, 92], "9º BPM": [108, 59, 92, 66],
    "10º BPM": [148, 47], "11º BPM": [202, 102, 109], "12º BPM": [232, 130],
    "13º BPM": [76, 38, 40], "14º BPM": [209, 148], "15º BPM": [144, 105],
    "16º BPM": [177, 136, 90], "17º BPM": [209, 172], "18º BPM": [219, 165],
    "19º BPM": [158, 149, 148], "20º BPM": [145, 134, 121], "21º BPM": [117, 243],
    "22º BPM": [140, 114], "23º BPM": [118, 82, 60, 66], "24º BPM": [180, 144],
    "25º BPM": [136, 81], "26º BPM": [142, 144], "27º BPM": [136, 76],
    "28º BPM": [101, 72], "29º BPM": [146, 94], "30º BPM": [117, 70],
    "31º BPM": [110, 51], "32º BPM": [155, 67], "33º BPM": [90, 35], "34º BPM": [120, 71],
}

TERRITORIES = {
    "1º BPM": ["Russas", "Limoeiro do Norte"], "2º BPM": ["Juazeiro do Norte", "Barbalha"],
    "3º BPM": ["Sobral", "Forquilha"], "4º BPM": ["Canindé", "Boa Viagem"],
    "5º BPM": ["Centro", "Carlito Pamplona"], "6º BPM": ["Parangaba", "Bairro de Fátima"],
    "7º BPM": ["Crateús", "Nova Russas"], "8º BPM": ["Aldeota", "Vicente Pinzón"],
    "9º BPM": ["Quixadá", "Senador Pompeu"], "10º BPM": ["Iguatu", "Jucás"],
    "11º BPM": ["Itapipoca", "Pentecoste"], "12º BPM": ["Centro", "Cumbuco"],
    "13º BPM": ["Tauá", "Parambu"], "14º BPM": ["Maracanaú"],
    "15º BPM": ["Eusébio", "Cascavel"], "16º BPM": ["Messejana", "Jangurussu"],
    "17º BPM": ["Conjunto Ceará", "Bom Jardim"], "18º BPM": ["Antônio Bezerra", "Parquelândia"],
    "19º BPM": ["Cambeba", "Aerolândia"], "20º BPM": ["Pirambu", "Barra do Ceará"],
    "21º BPM": ["Conjunto Esperança", "Maraponga"], "22º BPM": ["Papicu", "Dionísio Torres"],
    "23º BPM": ["Paraipaba", "São Gonçalo do Amarante"], "24º BPM": ["Maranguape", "Pacatuba"],
    "25º BPM": ["Horizonte", "Pacajus"], "26º BPM": ["Jurema", "Nova Metrópole"],
    "27º BPM": ["Tianguá", "São Benedito"], "28º BPM": ["Camocim", "Granja"],
    "29º BPM": ["Baturité", "Guaramiranga"], "30º BPM": ["Aracati", "Beberibe"],
    "31º BPM": ["Jaguaribe", "Alto Santo"], "32º BPM": ["Brejo Santo", "Mauriti"],
    "33º BPM": ["Campos Sales", "Assaré"], "34º BPM": ["Icó", "Várzea Alegre"],
}

HEADQUARTERS = {
    "1º BPM": "Russas", "2º BPM": "Juazeiro do Norte", "3º BPM": "Sobral", "4º BPM": "Canindé",
    "5º BPM": "Fortaleza", "6º BPM": "Fortaleza", "7º BPM": "Crateús", "8º BPM": "Fortaleza",
    "9º BPM": "Quixadá", "10º BPM": "Iguatu", "11º BPM": "Itapipoca", "12º BPM": "Caucaia",
    "13º BPM": "Tauá", "14º BPM": "Maracanaú", "15º BPM": "Eusébio", "16º BPM": "Fortaleza",
    "17º BPM": "Fortaleza", "18º BPM": "Fortaleza", "19º BPM": "Fortaleza", "20º BPM": "Fortaleza",
    "21º BPM": "Fortaleza", "22º BPM": "Fortaleza", "23º BPM": "Paraipaba", "24º BPM": "Maranguape",
    "25º BPM": "Horizonte", "26º BPM": "Caucaia", "27º BPM": "Tianguá", "28º BPM": "Camocim",
    "29º BPM": "Baturité", "30º BPM": "Aracati", "31º BPM": "Jaguaribe", "32º BPM": "Brejo Santo",
    "33º BPM": "Campos Sales", "34º BPM": "Icó",
}

CAPITAL_ANCHORS = {
    "5º BPM": "Centro", "6º BPM": "Parangaba", "8º BPM": "Aldeota", "16º BPM": "Messejana",
    "17º BPM": "Conjunto Ceará I", "18º BPM": "Antônio Bezerra", "19º BPM": "Cambeba",
    "20º BPM": "Pirambu", "21º BPM": "Conjunto Esperança", "22º BPM": "Papicu",
}

CAUCAIA_ANCHORS = {
    "12º BPM": [-38.6566, -3.7362],
    "26º BPM": [-38.6358, -3.8355],
}

CRPM_BY_BPM = {
    "1º BPM": "8º CRPM", "2º BPM": "4º CRPM", "3º BPM": "3º CRPM", "4º BPM": "7º CRPM",
    "5º BPM": "1º CRPM", "6º BPM": "1º CRPM", "7º BPM": "3º CRPM", "8º BPM": "5º CRPM",
    "9º BPM": "8º CRPM", "10º BPM": "4º CRPM", "11º BPM": "7º CRPM", "12º BPM": "2º CRPM",
    "13º BPM": "4º CRPM", "14º BPM": "6º CRPM", "15º BPM": "6º CRPM", "16º BPM": "5º CRPM",
    "17º BPM": "1º CRPM", "18º BPM": "1º CRPM", "19º BPM": "5º CRPM", "20º BPM": "1º CRPM",
    "21º BPM": "1º CRPM", "22º BPM": "5º CRPM", "23º BPM": "2º CRPM", "24º BPM": "6º CRPM",
    "25º BPM": "6º CRPM", "26º BPM": "2º CRPM", "27º BPM": "3º CRPM", "28º BPM": "3º CRPM",
    "29º BPM": "7º CRPM", "30º BPM": "8º CRPM", "31º BPM": "8º CRPM", "32º BPM": "4º CRPM",
    "33º BPM": "4º CRPM", "34º BPM": "4º CRPM",
}

FORTALEZA_NEIGHBORHOOD_ALIASES = {
    "tauape": "saojoaodotauape",
    "sapirangacoite": "sapiranga",
    "boavistacastelao": "boavista",
}


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]", "", text.casefold())


def clean_unit(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace(" º", "º").strip())


def distribution_source() -> Path:
    source = next((path for path in DISTRI_VTR_CANDIDATES if path.exists()), None)
    if source is None:
        raise FileNotFoundError("A planilha DISTRI VTR (1).xlsx não foi encontrada nas fontes do projeto nem em Downloads.")
    return source


def read_distribution() -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]]]:
    workbook = load_workbook(distribution_source(), read_only=True, data_only=True)
    sheet = workbook["BASE"]
    rows = {}
    fortaleza_by_ais = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        assignment = {
            "municipio": str(row[0]).strip(),
            "risp": str(row[1]).strip(),
            "ais": str(row[2]).strip(),
            "crpm": clean_unit(row[3]),
            "bpm": clean_unit(row[4]),
        }
        municipality_key = normalized(row[0])
        rows[municipality_key] = assignment
        if municipality_key == "fortaleza":
            fortaleza_by_ais[assignment["ais"]] = {
                "ais": assignment["ais"],
                "crpm": assignment["crpm"],
                "bpm": assignment["bpm"],
            }

    neighborhood_assignments = {}
    for row in workbook["BAIRROS FORTALEZA"].iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        ais = str(row[0]).strip()
        assignment = fortaleza_by_ais.get(ais)
        if assignment is None:
            raise ValueError(f"Bairro de Fortaleza associado a uma AIS sem vínculo territorial: {row[1]} · {ais}")
        neighborhood_assignments[normalized(row[1])] = assignment

    workbook.close()
    assert len(rows) == 184, f"Esperados 184 municípios na distribuição; encontrados {len(rows)}"
    assert len(neighborhood_assignments) == 121, (
        f"Esperados 121 bairros na distribuição de Fortaleza; encontrados {len(neighborhood_assignments)}"
    )
    rows["fortaleza"] = {
        "municipio": "Fortaleza",
        "risp": "RISP Capital Oeste e Leste",
        "ais": "10 AIS territoriais",
        "crpm": "1º e 5º CRPM",
        "bpm": "10 BPMs territoriais",
    }
    return rows, neighborhood_assignments


def ensure_cartographic_sources() -> Path:
    shapefiles = list(IPECE_DIR.glob("*.shp"))
    if not shapefiles:
        IPECE_ARCHIVE.parent.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(IPECE_URL, IPECE_ARCHIVE)
        IPECE_DIR.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(IPECE_ARCHIVE) as archive:
            archive.extractall(IPECE_DIR)
        shapefiles = list(IPECE_DIR.glob("*.shp"))
    if not FORTALEZA_SOURCE.exists():
        MAP_DATA.mkdir(parents=True, exist_ok=True)
        urllib.request.urlretrieve(FORTALEZA_URL, FORTALEZA_SOURCE)
    if len(shapefiles) != 1:
        raise RuntimeError("Não foi possível identificar uma única malha municipal do IPECE.")
    return shapefiles[0]


def build_municipalities(distribution: dict[str, dict[str, str]], shapefile_path: Path):
    reader = shapefile.Reader(str(shapefile_path), encoding="utf-8")
    transformer = Transformer.from_crs("EPSG:31984", "EPSG:4326", always_xy=True)
    features = []
    centers = {}
    unmatched = []
    for record in reader.iterShapeRecords():
        properties = record.record.as_dict()
        municipality = properties["Municipio"].strip()
        assignment = distribution.get(normalized(municipality))
        if not assignment:
            unmatched.append(municipality)
            continue
        geometry = transform(transformer.transform, shape(record.shape.__geo_interface__))
        geometry = geometry.simplify(0.00035, preserve_topology=True)
        point = geometry.representative_point()
        centers[normalized(municipality)] = [round(point.x, 6), round(point.y, 6)]
        features.append({
            "type": "Feature",
            "properties": {
                "municipio": municipality,
                "codigo_ibge": properties["codigo_ibg"],
                "regiao_planejamento": properties["regiao"],
                "area_km2": round(float(properties["area_km2"]), 2),
                **assignment,
            },
            "geometry": mapping(geometry),
        })
    assert not unmatched, f"Municípios sem vínculo territorial: {unmatched}"
    assert len(features) == 184
    collection = {
        "type": "FeatureCollection",
        "name": "Limites municipais do Ceará 2026",
        "source": "IPECE — Projeto Atlas de Limites",
        "features": features,
    }
    output = MAP_DATA / "municipios-ceara-2026.geojson"
    output.write_text(json.dumps(collection, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return centers, collection


def add_neighborhood_assignments(source: dict, neighborhood_assignments: dict[str, dict[str, str]]) -> dict:
    features = []
    unmatched = []
    for feature in source["features"]:
        properties = feature["properties"]
        name = str(properties.get("Nome") or properties.get("bairro") or "").strip()
        source_key = FORTALEZA_NEIGHBORHOOD_ALIASES.get(normalized(name), normalized(name))
        assignment = neighborhood_assignments.get(source_key)
        if assignment is None:
            unmatched.append(name)
            continue
        features.append({
            "type": "Feature",
            "properties": {"bairro": name, **assignment},
            "geometry": feature["geometry"],
        })
    assert not unmatched, f"Bairros de Fortaleza sem vínculo de CRPM: {unmatched}"
    assert len(features) == 121
    return {"type": "FeatureCollection", "features": features}


def build_fortaleza_neighborhoods(neighborhood_assignments: dict[str, dict[str, str]]) -> tuple[dict[str, list[float]], dict]:
    source = json.loads(FORTALEZA_SOURCE.read_text(encoding="utf-8"))
    features = []
    centers = {}
    for feature in source["features"]:
        name = feature["properties"]["Nome"].strip()
        geometry = shape(feature["geometry"]).simplify(0.00008, preserve_topology=True)
        point = geometry.representative_point()
        centers[normalized(name)] = [round(point.x, 6), round(point.y, 6)]
        features.append({
            "type": "Feature",
            "properties": {"bairro": name},
            "geometry": mapping(geometry),
        })
    assert len(features) == 121
    collection = add_neighborhood_assignments(
        {"type": "FeatureCollection", "features": features},
        neighborhood_assignments,
    )
    output = MAP_DATA / "bairros-fortaleza.geojson"
    output.write_text(json.dumps(collection, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return centers, collection


def build_crpm_regions(municipalities: dict, neighborhoods: dict) -> dict:
    grouped_geometries = {f"{number}º CRPM": [] for number in range(1, 9)}
    municipality_counts = {crpm: 0 for crpm in grouped_geometries}
    neighborhood_counts = {crpm: 0 for crpm in grouped_geometries}

    for feature in municipalities["features"]:
        properties = feature["properties"]
        if properties["municipio"] == "Fortaleza":
            continue
        crpm = properties["crpm"]
        grouped_geometries[crpm].append(shape(feature["geometry"]))
        municipality_counts[crpm] += 1

    for feature in neighborhoods["features"]:
        crpm = feature["properties"]["crpm"]
        grouped_geometries[crpm].append(shape(feature["geometry"]))
        neighborhood_counts[crpm] += 1

    features = []
    for crpm, geometries in grouped_geometries.items():
        assert geometries, f"Nenhuma geometria encontrada para {crpm}"
        region = unary_union(geometries).buffer(0)
        region_parts = list(region.geoms) if isinstance(region, MultiPolygon) else [region]
        cleaned_parts = []
        for part in region_parts:
            meaningful_holes = [
                ring.coords
                for ring in part.interiors
                if Polygon(ring).area >= 0.0001
            ]
            cleaned_parts.append(Polygon(part.exterior.coords, meaningful_holes))
        region = MultiPolygon(cleaned_parts) if len(cleaned_parts) > 1 else cleaned_parts[0]
        label_point = region.representative_point()
        battalions = [name for name, regional_command in CRPM_BY_BPM.items() if regional_command == crpm]
        features.append({
            "type": "Feature",
            "properties": {
                "crpm": crpm,
                "batalhoes": battalions,
                "quantidade_batalhoes": len(battalions),
                "municipios_integrais": municipality_counts[crpm],
                "bairros_fortaleza": neighborhood_counts[crpm],
                "label_coordinates": [round(label_point.x, 6), round(label_point.y, 6)],
            },
            "geometry": mapping(region),
        })

    collection = {
        "type": "FeatureCollection",
        "name": "Divisões territoriais dos oito CRPMs",
        "source": "DISTRI VTR (1).xlsx · limites IPECE 2026 · bairros IPLANFOR 2023",
        "features": features,
    }
    output = MAP_DATA / "crpm-regioes.geojson"
    output.write_text(json.dumps(collection, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return collection


def read_situational_rows() -> dict[str, dict[str, object]]:
    result = {}
    with CSV_SOURCE.open(encoding="utf-8", newline="") as stream:
        for row in csv.DictReader(stream):
            name = row["batalhao"]
            if name == "TOTAL":
                continue
            result[name] = {
                "exoneracoes": int(row["exoneracoes"]),
                "demissoes": int(row["demissoes"]),
                "requeridas_2025": int(row["requeridas_2025"]),
                "requeridas_2026": int(row["requeridas_2026"]),
                "requeridas": int(row["requeridas_total"]),
                "movimentacoes": int(row["movimentacoes_saldo"]),
                "situacao": int(row["situacao_media"]),
                "perdas": int(row["deficit_situacional"]),
                "reestruturacao": int(row["reestruturacao_interior_litoral"] or 0),
                "necessidade": int(row["necessidade_de_efetivo"]),
            }
    assert len(result) == 34
    return result


def build_battalions(municipal_centers, neighborhood_centers, distribution):
    situational = read_situational_rows()
    assignments_by_bpm = {}
    for assignment in distribution.values():
        assignments_by_bpm.setdefault(assignment["bpm"], []).append(assignment["municipio"])
    features = []
    for name, values in situational.items():
        headquarters = HEADQUARTERS[name]
        if name in CAPITAL_ANCHORS:
            anchor_name = CAPITAL_ANCHORS[name]
            coordinates = neighborhood_centers[normalized(anchor_name)]
            anchor_kind = "bairro oficial de Fortaleza"
        elif name in CAUCAIA_ANCHORS:
            coordinates = CAUCAIA_ANCHORS[name]
            anchor_name = TERRITORIES[name][0]
            anchor_kind = "bairro de referência em Caucaia"
        else:
            coordinates = municipal_centers[normalized(headquarters)]
            anchor_name = headquarters
            anchor_kind = "município-sede"
        municipality_coverage = sorted(assignments_by_bpm.get(name, []))
        if headquarters in {"Fortaleza", "Caucaia"}:
            municipality_coverage = [headquarters]
        effective = sum(COMPANY_STRENGTHS[name])
        features.append({
            "type": "Feature",
            "properties": {
                "batalhao": name,
                "numero": int(re.match(r"\d+", name).group()),
                "sede": headquarters,
                "ancora": anchor_name,
                "tipo_ancora": anchor_kind,
                "localidades_referencia": TERRITORIES[name],
                "municipios_cobertos": municipality_coverage,
                "quantidade_municipios": len(municipality_coverage),
                "crpm": CRPM_BY_BPM[name],
                "efetivo": effective,
                "companhias": len(COMPANY_STRENGTHS[name]),
                **values,
            },
            "geometry": {"type": "Point", "coordinates": coordinates},
        })
    assert len(features) == 34
    output = MAP_DATA / "batalhoes-situacao.geojson"
    output.write_text(json.dumps({"type": "FeatureCollection", "features": features}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def main():
    MAP_DATA.mkdir(parents=True, exist_ok=True)
    shapefile_path = ensure_cartographic_sources()
    distribution, neighborhood_assignments = read_distribution()
    municipal_centers, municipalities = build_municipalities(distribution, shapefile_path)
    neighborhood_centers, neighborhoods = build_fortaleza_neighborhoods(neighborhood_assignments)
    build_crpm_regions(municipalities, neighborhoods)
    build_battalions(municipal_centers, neighborhood_centers, distribution)
    print("Mapa validado: 8 CRPMs, 184 municípios, 121 bairros de Fortaleza e 34 BPMs.")


if __name__ == "__main__":
    main()
